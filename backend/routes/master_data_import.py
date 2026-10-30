from flask import Blueprint, request, jsonify
from flask_jwt_extended import jwt_required
from utils.auth_decorators import require_permission
from models import db
from models.finance import Account
from models.product import Product, Material
from utils.uom_helpers import upsert_item_uom_conversions_from_row
import openpyxl
import io

master_data_import_bp = Blueprint('master_data_import', __name__)

CATEGORY_MAP = {
    'BARANG JADI': ('product', 'finished_goods'),
    'PREMIX': ('product', 'premix'),
    'BARANG DALAM PENYELESAIAN': ('product', 'wip'),
    'WASTE': ('product', 'waste'),
    'BAHAN BAKU': ('material', 'raw_materials'),
    'BAHAN KIMIA': ('material', 'chemical_materials'),
    'BAHAN PEMBANTU': ('material', 'packaging_materials'),
    # Item biaya/jasa Accurate (jenis NON/SVC/PDC) -> barang non persediaan (tanpa stok, langsung ke akun beban)
    'Umum': ('material', 'non_inventory'),
    'LAIN-LAIN': ('material', 'non_inventory'),
    # jenis INV/GROUP di kategori LAIN-LAIN/Umum dilewati (lihat NON_INVENTORY_JENIS)
}

# Column indices (0-based), matching the standard Accurate "Barang & Jasa" export template
COL = dict(
    kategori=1, kode=2, nama=3, jenis=4, satuan=5,
    satuan_2=6, rasio_satuan_2=7, satuan_3=8, rasio_satuan_3=9,
    harga_beli=32, min_beli=33, ppn=35,
    gl_persediaan=47, gl_hpp=48, gl_barang_terkirim=49, gl_pembelian_belum_tertagih=50,
    gl_penjualan=51, gl_diskon_penjualan=52, gl_retur_penjualan=53, gl_retur_pembelian=54,
    kelompok=56, erp_approval=57,
    lead_time=71, self_life_retest=72,
    non_aktif=78,
)

NON_INVENTORY_JENIS = ('NON', 'SVC', 'PDC')

AKUN_KEYS = ('gl_persediaan', 'gl_hpp', 'gl_barang_terkirim', 'gl_pembelian_belum_tertagih',
             'gl_penjualan', 'gl_diskon_penjualan', 'gl_retur_penjualan', 'gl_retur_pembelian')


@master_data_import_bp.route('/import-items', methods=['POST'])
@jwt_required()
@require_permission('products.create')
def import_items():
    """Bulk import Product/Material rows from the standard Accurate 'Barang & Jasa'
    Excel export - same routing/logic as the one-off script used to seed erp_db_v2,
    now reusable from the Master Data page's Import Excel button."""
    if 'file' not in request.files:
        return jsonify({'error': 'File tidak ditemukan'}), 400

    file = request.files['file']
    if not file.filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'File harus berformat .xlsx atau .xls'}), 400

    sheet_name = request.form.get('sheet_name', 'Barang & Jasa (SORTIR)')

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file.read()), data_only=True)
    except Exception as e:
        return jsonify({'error': f'Gagal membaca file Excel: {e}'}), 400

    if sheet_name not in wb.sheetnames:
        return jsonify({
            'error': f"Sheet '{sheet_name}' tidak ditemukan. Sheet tersedia: {', '.join(wb.sheetnames)}"
        }), 400

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))[1:]
    return jsonify(import_item_rows(rows)), 200


def import_item_rows(rows, commit=True):
    """Create Product/Material rows from 'Barang & Jasa' export rows (header already removed).
    Each row runs in its own savepoint, so one bad row never undoes the others. Returns a summary dict."""
    accounts = {a.account_code: a.id for a in Account.query.all()}

    created_product, created_material, skipped_category, skipped_exists = 0, 0, 0, 0
    missing_gl_codes = set()
    errors = []

    for r in rows:
        try:
            with db.session.begin_nested():
                if len(r) <= max(COL.values()):
                    continue
                kategori = (r[COL['kategori']] or '').strip()
                route = CATEGORY_MAP.get(kategori)
                if not route:
                    skipped_category += 1
                    continue
                table, mat_type = route
                jenis = str(r[COL['jenis']] or '').strip().upper()
                if mat_type == 'non_inventory' and jenis not in NON_INVENTORY_JENIS:
                    skipped_category += 1   # barang persediaan nyasar di kategori Umum, atau barang tipe grup
                    continue

                code = str(r[COL['kode']]).strip()
                name = (r[COL['nama']] or '').strip()
                uom = (r[COL['satuan']] or ('UNIT' if mat_type == 'non_inventory' else 'PCS')).strip() or ('UNIT' if mat_type == 'non_inventory' else 'PCS')
                harga_beli = r[COL['harga_beli']] or 0
                min_beli = r[COL['min_beli']]
                # 0 in this column means "not actually set" in practice (743/745
                # rows in the real source file are a template default of 0, only
                # 2 have a genuine value) - store as None, not a real MOQ of 0.
                min_beli = float(min_beli) if isinstance(min_beli, (int, float)) and min_beli > 0 else None
                ppn_raw = r[COL['ppn']]
                ppn_code = ppn_raw.strip().upper() if isinstance(ppn_raw, str) and ppn_raw.strip() and ppn_raw.strip().upper() in ('Y', 'A', 'B', 'L') else None
                kelompok = r[COL['kelompok']]
                kelompok = kelompok.strip() if isinstance(kelompok, str) and kelompok.strip() else None
                erp_approval_raw = r[COL['erp_approval']]
                erp_approval = True if not erp_approval_raw else str(erp_approval_raw).strip().upper() not in ('TIDAK', 'FALSE', 'NO', '0')
                lead_time = r[COL['lead_time']]
                lead_time = int(lead_time) if isinstance(lead_time, (int, float)) else None
                self_life_retest = r[COL['self_life_retest']]
                self_life_retest = int(self_life_retest) if isinstance(self_life_retest, (int, float)) else None
                non_aktif = (r[COL['non_aktif']] or '').strip().upper()
                is_active = non_aktif != 'YA'

                akun_ids = {}
                for key in AKUN_KEYS:
                    raw = r[COL[key]]
                    if raw is None:
                        akun_ids[key] = None
                        continue
                    code_str = str(raw).strip()
                    acc_id = accounts.get(code_str)
                    akun_ids[key] = acc_id
                    if acc_id is None:
                        missing_gl_codes.add(code_str)

                akun_kwargs = dict(
                    akun_persediaan_id=akun_ids['gl_persediaan'],
                    akun_hpp_id=akun_ids['gl_hpp'],
                    akun_barang_terkirim_id=akun_ids['gl_barang_terkirim'],
                    akun_pembelian_belum_tertagih_id=akun_ids['gl_pembelian_belum_tertagih'],
                    akun_penjualan_id=akun_ids['gl_penjualan'],
                    akun_diskon_penjualan_id=akun_ids['gl_diskon_penjualan'],
                    akun_retur_penjualan_id=akun_ids['gl_retur_penjualan'],
                    akun_retur_pembelian_id=akun_ids['gl_retur_pembelian'],
                )

                if mat_type == 'non_inventory':
                    # Pada item biaya/jasa Accurate, kolom "Persediaan" memuat akun BEBAN-nya; akun lain tidak relevan
                    # kecuali Penjualan pada item pendapatan (PDC).
                    akun_kwargs = dict(akun_beban_id=akun_ids['gl_persediaan'],
                                       akun_penjualan_id=akun_ids['gl_penjualan'] if jenis == 'PDC' else None)

                if table == 'product':
                    if Product.query.filter_by(code=code).first():
                        skipped_exists += 1
                        continue
                    obj = Product(
                        code=code, name=name, primary_uom=uom,
                        price=0, cost=harga_beli, material_type=mat_type,
                        kelompok=kelompok, erp_approval=erp_approval, ppn_code=ppn_code,
                        lead_time_days=lead_time or 0,
                        min_order_qty=min_beli,
                        is_active=is_active,
                        **akun_kwargs,
                    )
                    if mat_type == 'wip':
                        obj.retest_period_days = self_life_retest
                    else:
                        obj.self_life_days = self_life_retest
                    db.session.add(obj)
                    db.session.flush()
                    upsert_item_uom_conversions_from_row(
                        uom, r[COL['satuan_2']], r[COL['rasio_satuan_2']],
                        r[COL['satuan_3']], r[COL['rasio_satuan_3']], product_id=obj.id,
                    )
                    created_product += 1
                else:
                    if Material.query.filter_by(code=code).first():
                        skipped_exists += 1
                        continue
                    obj = Material(
                        code=code, name=name, material_type=mat_type,
                        category='Non Persediaan' if mat_type == 'non_inventory' else kategori,
                        is_excluded_from_mrp=(mat_type == 'non_inventory'),
                        primary_uom=uom, cost_per_unit=harga_beli,
                        kelompok=kelompok, erp_approval=erp_approval, ppn_code=ppn_code,
                        lead_time_days=lead_time or 0,
                        min_order_qty=min_beli,
                        expiry_days=self_life_retest,
                        is_active=is_active,
                        **akun_kwargs,
                    )
                    db.session.add(obj)
                    db.session.flush()
                    upsert_item_uom_conversions_from_row(
                        uom, r[COL['satuan_2']], r[COL['rasio_satuan_2']],
                        r[COL['satuan_3']], r[COL['rasio_satuan_3']], material_id=obj.id,
                    )
                    created_material += 1

                db.session.flush()
        except Exception as e:
            errors.append(f"{r[COL['kode']] if len(r) > COL['kode'] else '?'}: {e}")

    if commit:
        db.session.commit()
    else:
        db.session.flush()

    return {
        'created_product': created_product,
        'created_material': created_material,
        'skipped_category': skipped_category,
        'skipped_exists': skipped_exists,
        'errors': errors[:50],
        'error_count': len(errors),
        'missing_gl_codes': sorted(missing_gl_codes)[:50],
        'missing_gl_code_count': len(missing_gl_codes),
    }
