"""
HR Reports (2026-09-28) - backend untuk frontend/src/pages/HR/Reports.tsx.

Ditemukan saat audit HR module maturation: halaman ini sudah lama ada di
frontend (memanggil GET /api/hr/reports/generate dan /export) tapi endpoint-
nya TIDAK PERNAH dibuat di backend sama sekali - 404 di setiap pemakaian,
dan halaman ini juga tidak ditautkan dari sidebar manapun (orphan ganda).
Dibangun dari nol di sini, mengikuti kontrak yang sudah dipakai frontend-nya
(bukan didesain ulang) supaya frontend yang sudah ada langsung berfungsi
tanpa perlu diubah:
  GET /generate?report_type=...&date_from&date_to&department
      -> {"success": true, "data": {"headers": [...], "rows": [[...], ...]}}
  GET /export?report_type=...&format=excel|pdf&date_from&date_to&department
      -> file (xlsx atau pdf)

5 report_type sesuai REPORT_TYPES di Reports.tsx: employee_list,
attendance_summary, leave_report, payroll_summary, department_headcount.
"""
import io
from datetime import datetime

from flask import Blueprint, request, jsonify, send_file
from flask_jwt_extended import jwt_required
from utils.auth_decorators import require_permission

from models import db
from models.hr import Employee, Department, Attendance, StaffLeaveRequest
from models.hr_extended import PayrollRecord, PayrollPeriod

hr_reports_bp = Blueprint('hr_reports', __name__)


def _parse_date(value):
    if not value:
        return None
    try:
        return datetime.strptime(value, '%Y-%m-%d').date()
    except ValueError:
        return None


def _build_report(report_type, date_from, date_to, department_code):
    """Return (headers: list[str], rows: list[list]) untuk satu report_type.
    Sengaja query sederhana per tipe (bukan satu query generik) - tiap
    laporan punya bentuk & join yang beda, memaksa satu abstraksi generik
    di sini cuma akan bikin kode lebih susah dibaca tanpa manfaat nyata
    untuk cuma 5 tipe laporan."""

    if report_type == 'employee_list':
        query = Employee.query.filter_by(is_active=True)
        if department_code:
            query = query.join(Department).filter(Department.code == department_code)
        employees = query.order_by(Employee.full_name).all()
        headers = ['No', 'Nama', 'NIK', 'Departemen', 'Jabatan', 'Status', 'Tanggal Masuk']
        rows = [
            [
                i + 1,
                e.full_name,
                e.nik or '-',
                e.department.name if e.department else '-',
                e.position or '-',
                e.status,
                e.hire_date.isoformat() if e.hire_date else '-',
            ]
            for i, e in enumerate(employees)
        ]
        return headers, rows

    if report_type == 'attendance_summary':
        query = db.session.query(Employee).filter(Employee.is_active.is_(True))
        if department_code:
            query = query.join(Department).filter(Department.code == department_code)
        employees = query.order_by(Employee.full_name).all()

        headers = ['Nama', 'Hadir', 'Terlambat', 'Absen', 'Total Jam Kerja']
        rows = []
        for e in employees:
            att_query = Attendance.query.filter(Attendance.employee_id == e.id)
            if date_from:
                att_query = att_query.filter(Attendance.attendance_date >= date_from)
            if date_to:
                att_query = att_query.filter(Attendance.attendance_date <= date_to)
            records = att_query.all()
            if not records:
                continue
            present = sum(1 for r in records if r.status == 'present')
            late = sum(1 for r in records if r.status == 'late')
            absent = sum(1 for r in records if r.status == 'absent')
            total_hours = sum(float(r.worked_hours or 0) for r in records)
            rows.append([e.full_name, present, late, absent, round(total_hours, 2)])
        return headers, rows

    if report_type == 'leave_report':
        # StaffLeaveRequest adalah satu-satunya sistem cuti sejak konsolidasi
        # (lihat commit "consolidate two parallel leave systems").
        query = StaffLeaveRequest.query
        if date_from:
            query = query.filter(StaffLeaveRequest.start_date >= date_from)
        if date_to:
            query = query.filter(StaffLeaveRequest.start_date <= date_to)
        if department_code:
            query = query.join(Employee, StaffLeaveRequest.employee_id == Employee.id).join(Department).filter(Department.code == department_code)
        requests_ = query.order_by(StaffLeaveRequest.start_date.desc()).all()

        headers = ['Nama', 'Tipe', 'Tanggal Mulai', 'Tanggal Selesai', 'Jumlah Hari', 'Status']
        rows = [
            [
                r.staff_name,
                r.leave_type,
                r.start_date.isoformat() if r.start_date else '-',
                r.end_date.isoformat() if r.end_date else '-',
                r.total_days,
                r.status,
            ]
            for r in requests_
        ]
        return headers, rows

    if report_type == 'payroll_summary':
        query = db.session.query(PayrollRecord).join(PayrollPeriod)
        if date_from:
            query = query.filter(PayrollPeriod.end_date >= date_from)
        if date_to:
            query = query.filter(PayrollPeriod.start_date <= date_to)
        if department_code:
            query = query.join(Employee, PayrollRecord.employee_id == Employee.id).join(Department).filter(Department.code == department_code)
        records = query.order_by(PayrollPeriod.start_date.desc()).all()

        headers = ['Nama', 'Periode', 'Gaji Pokok', 'Tunjangan', 'Lembur', 'Potongan', 'Gaji Bersih']
        rows = [
            [
                r.employee.full_name if r.employee else '-',
                r.period.period_name if r.period else '-',
                float(r.basic_salary or 0),
                float(r.allowances or 0),
                float(r.overtime_amount or 0),
                float(r.total_deductions or 0),
                float(r.net_salary or 0),
            ]
            for r in records
        ]
        return headers, rows

    if report_type == 'department_headcount':
        # Filter departemen sengaja diabaikan untuk laporan ini - tujuannya
        # justru membandingkan SEMUA departemen sekaligus, memfilter ke satu
        # departemen akan membuat laporan ini tidak berguna.
        departments = Department.query.filter_by(is_active=True).order_by(Department.name).all()
        headers = ['Departemen', 'Jumlah Karyawan Aktif']
        rows = [
            [d.name, Employee.query.filter_by(department_id=d.id, is_active=True).count()]
            for d in departments
        ]
        return headers, rows

    return None, None


@hr_reports_bp.route('/generate', methods=['GET'])
@jwt_required()
@require_permission('hr.view')
def generate_report():
    report_type = request.args.get('report_type')
    date_from = _parse_date(request.args.get('date_from'))
    date_to = _parse_date(request.args.get('date_to'))
    department_code = request.args.get('department') or None

    headers, rows = _build_report(report_type, date_from, date_to, department_code)
    if headers is None:
        return jsonify({'error': f'report_type tidak dikenal: {report_type}'}), 400

    return jsonify({'success': True, 'data': {'headers': headers, 'rows': rows}}), 200


@hr_reports_bp.route('/export', methods=['GET'])
@jwt_required()
@require_permission('hr.view')
def export_report():
    report_type = request.args.get('report_type')
    fmt = request.args.get('format', 'excel')
    date_from = _parse_date(request.args.get('date_from'))
    date_to = _parse_date(request.args.get('date_to'))
    department_code = request.args.get('department') or None

    headers, rows = _build_report(report_type, date_from, date_to, department_code)
    if headers is None:
        return jsonify({'error': f'report_type tidak dikenal: {report_type}'}), 400

    if fmt == 'pdf':
        return _export_pdf(report_type, headers, rows)
    return _export_excel(report_type, headers, rows)


def _export_excel(report_type, headers, rows):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = report_type[:31]  # Excel sheet name max 31 karakter

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='059669')

    for row in rows:
        ws.append(row)

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f'hr_report_{report_type}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def _export_pdf(report_type, headers, rows):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()

    title = report_type.replace('_', ' ').title()
    elements = [Paragraph(f'Laporan HR - {title}', styles['Title']), Spacer(1, 12)]

    table_data = [headers] + [[str(c) for c in row] for row in rows]
    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#059669')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F4F6')]),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)

    return send_file(buffer, as_attachment=True, download_name=f'hr_report_{report_type}.pdf', mimetype='application/pdf')
