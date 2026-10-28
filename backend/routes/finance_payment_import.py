"""Impor transaksi Penerimaan / Pembayaran dari Excel (.xlsx) atau CSV.

- source=kas      : transaksi Kas/Bank mandiri (akun lawan). Kolom: Tanggal, No. Bukti, Kas/Bank, Akun Lawan, Nilai, Pihak, Keterangan, No. Referensi
- source=invoice  : pelunasan faktur. Kolom: Tanggal, No. Bukti, Kas/Bank, No. Faktur, Nilai, Keterangan, No. Referensi
Baris dengan No. Bukti yang sama digabung jadi satu transaksi (rincian ganda). Nilai boleh negatif, tidak boleh 0.
Pratinjau (dry run) memvalidasi dengan aturan yang sama dengan form lalu membatalkan semuanya; simpan sungguhan per transaksi.
"""
import csv
import io
from datetime import date, datetime

from flask import Blueprint, jsonify, request, send_file
from flask_jwt_extended import get_jwt_identity, jwt_required

from models import db
from models.finance import Account, Invoice
from utils.auth_decorators import require_permission

finance_payment_import_bp = Blueprint('finance_payment_import', __name__)

COLUMNS = {
    'kas': [('tanggal', 'Tanggal'), ('no_bukti', 'No. Bukti'), ('kas_bank', 'Kas/Bank'), ('akun_lawan', 'Akun Lawan'),
            ('nilai', 'Nilai'), ('pihak', 'Pihak'), ('keterangan', 'Keterangan'), ('referensi', 'No. Referensi')],
    'invoice': [('tanggal', 'Tanggal'), ('no_bukti', 'No. Bukti'), ('kas_bank', 'Kas/Bank'), ('no_faktur', 'No. Faktur'),
                ('nilai', 'Nilai'), ('keterangan', 'Keterangan'), ('referensi', 'No. Referensi')],
}
ALIASES = {
    'tanggal': ['tanggal', 'tgl', 'date'], 'no_bukti': ['no. bukti', 'no bukti', 'nomor bukti', 'no. transaksi', 'no transaksi'],
    'kas_bank': ['kas/bank', 'kas bank', 'kas', 'bank', 'akun kas/bank'], 'akun_lawan': ['akun lawan', 'akun', 'akun perkiraan'],
    'nilai': ['nilai', 'jumlah', 'amount'], 'pihak': ['pihak', 'dibayarkan kepada', 'diterima dari', 'dibayarkan kepada / diterima dari'],
    'keterangan': ['keterangan', 'deskripsi', 'catatan', 'notes'], 'referensi': ['no. referensi', 'no referensi', 'referensi', 'ref'],
    'no_faktur': ['no. faktur', 'no faktur', 'nomor faktur', 'faktur', 'invoice'],
}


def _kind(args):
    t = args.get('type')
    s = args.get('source')
    if t not in ('receipt', 'payment') or s not in ('kas', 'invoice'):
        raise ValueError('Parameter type (receipt/payment) dan source (kas/invoice) wajib diisi')
    return t, s


@finance_payment_import_bp.route('/template', methods=['GET'])
@jwt_required()
@require_permission('finance.view')
def template():
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    try:
        ptype, source = _kind(request.args)
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Transaksi'
    heads = [h for _, h in COLUMNS[source]]
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='047857')
        c.alignment = Alignment(horizontal='center')
    today = date.today().isoformat()
    if source == 'kas':
        ws.append([today, '', '100101', '700009', 1500000, 'PT Contoh', 'Contoh biaya admin bank', 'TRF-001'])
        ws.append([today, 'BUKTI-02', '100101', '600101', 2000000, 'CV Contoh', 'Contoh dibagi dua akun', ''])
        ws.append([today, 'BUKTI-02', '100101', '600102', -250000, '', 'Potongan (nilai negatif)', ''])
    else:
        ws.append([today, '', '100101', 'INV-202610-00001', 500000, 'Pelunasan sebagian', 'TRF-001'])
        ws.append([today, 'BUKTI-02', '100101', 'INV-202610-00002', 300000, 'Dua faktur satu bukti', ''])
        ws.append([today, 'BUKTI-02', '100101', 'INV-202610-00003', 200000, '', ''])
    for i, w in enumerate([14, 16, 12, 18, 16, 28, 34, 16], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    guide = wb.create_sheet('Petunjuk')
    rows = [
        'Petunjuk impor', '',
        'Baris pertama (judul kolom) jangan diubah. Hapus baris contoh sebelum mengisi data Anda.',
        'Tanggal: YYYY-MM-DD atau DD/MM/YYYY.',
        'No. Bukti: kosongkan agar nomor dibuat otomatis. Baris dengan No. Bukti sama digabung menjadi satu transaksi.',
        'Kas/Bank: kode akun kas/bank (contoh 100101).',
        'Nilai: angka tanpa simbol Rp. Boleh negatif (pengurang/pembalikan), tidak boleh 0.',
    ]
    rows += (['Akun Lawan: kode akun perkiraan (bukan akun induk, bukan akun piutang/hutang kontrol).', 'Pihak: nama yang menerima/membayar (opsional).']
             if source == 'kas' else
             ['No. Faktur: nomor faktur yang dilunasi. Semua faktur dalam satu No. Bukti harus milik pelanggan/pemasok yang sama.'])
    for r in rows:
        guide.append([r])
    guide['A1'].font = Font(bold=True, size=13)
    guide.column_dimensions['A'].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True,
                     download_name=f'template_import_{"penerimaan" if ptype == "receipt" else "pembayaran"}_{source}.xlsx')


def _read_rows(file_storage, source):
    """-> list of (excel_row_number, {field: value}) using the header row; raises ValueError on a bad file."""
    name = (file_storage.filename or '').lower()
    raw = file_storage.read()
    if name.endswith('.csv'):
        text = raw.decode('utf-8-sig', errors='replace')
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=',;\t') if text.strip() else csv.excel
        table = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    elif name.endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        ws = wb['Transaksi'] if 'Transaksi' in wb.sheetnames else wb.worksheets[0]
        table = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        raise ValueError('Format file harus .xlsx atau .csv')
    if not table:
        raise ValueError('File kosong')
    head = [str(h or '').strip().lower() for h in table[0]]
    index = {}
    for key, _label in COLUMNS[source]:
        for alias in ALIASES[key]:
            if alias in head:
                index[key] = head.index(alias)
                break
    missing = [label for key, label in COLUMNS[source] if key not in index and key in ('tanggal', 'kas_bank', 'nilai', 'akun_lawan', 'no_faktur')]
    if missing:
        raise ValueError('Kolom wajib tidak ditemukan: ' + ', '.join(missing) + '. Gunakan template dari tombol "Unduh Template".')
    out = []
    for n, row in enumerate(table[1:], start=2):
        rec = {k: (row[i] if i < len(row) else None) for k, i in index.items()}
        if all(v in (None, '') for v in rec.values()):
            continue
        out.append((n, rec))
    if not out:
        raise ValueError('Tidak ada baris data di file')
    if len(out) > 1000:
        raise ValueError('Maksimal 1000 baris per impor')
    return out


def _parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or '').strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    raise ValueError(f'Tanggal "{s}" tidak valid (pakai YYYY-MM-DD atau DD/MM/YYYY)')


def _parse_amount(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or '').strip().replace('Rp', '').replace(' ', '')
    if not s:
        raise ValueError('Nilai kosong')
    if ',' in s and '.' in s:                       # 1.234.567,89
        s = s.replace('.', '').replace(',', '.')
    elif ',' in s:
        s = s.replace(',', '.') if s.count(',') == 1 and len(s.split(',')[1]) != 3 else s.replace(',', '')
    elif s.count('.') > 1:
        s = s.replace('.', '')
    try:
        return float(s)
    except ValueError:
        raise ValueError(f'Nilai "{v}" bukan angka')


def build_transactions(rows, ptype, source):
    """Group rows -> [{'key','rows','errors','data'}]; 'data' is the create_payment_core payload when there are no errors."""
    groups, order = {}, []
    for n, rec in rows:
        key = str(rec.get('no_bukti') or '').strip() or f'__row{n}'
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append((n, rec))
    accounts = {}
    txs = []

    def account(code, want_cash=False):
        code = str(code or '').strip()
        if code.endswith('.0') and code[:-2].isdigit():
            code = code[:-2]
        if code not in accounts:
            accounts[code] = Account.query.filter_by(account_code=code).first()
        acc = accounts[code]
        if not acc:
            raise ValueError(f'Akun "{code}" tidak ditemukan')
        if acc.is_header:
            raise ValueError(f'Akun {code} adalah akun induk, pilih akun detail')
        if want_cash and not acc.is_cash_bank:
            raise ValueError(f'Akun {code} bukan akun Kas/Bank')
        return acc

    for key in order:
        items = groups[key]
        tx = {'key': key, 'number': None if key.startswith('__row') else key, 'rows': [n for n, _ in items], 'errors': [], 'data': None}
        first = items[0][1]
        zero_skipped = 0
        try:
            pdate = _parse_date(first.get('tanggal'))
            bank = account(first.get('kas_bank'), want_cash=True)
            if source == 'kas':
                lines = []
                for n, rec in items:
                    try:
                        amt = _parse_amount(rec.get('nilai'))
                        acc_line = account(rec.get('akun_lawan'))
                        if amt == 0:
                            zero_skipped += 1
                            continue
                        lines.append({'account_id': acc_line.id, 'amount': amt,
                                      'description': (str(rec.get('keterangan')).strip() if rec.get('keterangan') else None)})
                    except ValueError as e:
                        tx['errors'].append(f'Baris {n}: {e}')
                body = {'lines': lines}
            else:
                allocs = []
                for n, rec in items:
                    try:
                        num = str(rec.get('no_faktur') or '').strip()
                        inv = Invoice.query.filter_by(invoice_number=num).first()
                        if not inv:
                            raise ValueError(f'Faktur "{num}" tidak ditemukan')
                        if inv.invoice_type != ('sales' if ptype == 'receipt' else 'purchase'):
                            raise ValueError(f'Faktur {num} bukan faktur {"penjualan" if ptype == "receipt" else "pembelian"}')
                        amt = _parse_amount(rec.get('nilai'))
                        if amt == 0:
                            zero_skipped += 1
                            continue
                        allocs.append({'invoice_id': inv.id, 'amount': amt})
                    except ValueError as e:
                        tx['errors'].append(f'Baris {n}: {e}')
                body = {'allocations': allocs}
            if not tx['errors'] and not (body.get('lines') or body.get('allocations')):
                tx['errors'].append(f'Baris {tx["rows"][0]}: semua baris bernilai 0, tidak ada yang bisa disimpan')
            if not tx['errors']:
                tx['data'] = {
                    **body, 'payment_date': pdate.isoformat(), 'payment_type': ptype, 'payment_method': 'bank_transfer',
                    'bank_account_id': bank.id, 'payment_number': tx['number'],
                    'reference_number': str(first.get('referensi')).strip() if first.get('referensi') else None,
                    'notes': str(first.get('keterangan')).strip() if (source == 'invoice' and first.get('keterangan')) else None,
                    'paid_to': str(first.get('pihak')).strip() if (source == 'kas' and first.get('pihak')) else None,
                }
                tx['summary'] = {'date': pdate.isoformat(), 'bank': f'{bank.account_code} {bank.account_name}',
                                 'total': round(sum(float(i['amount']) for i in body.get('lines') or body.get('allocations')), 2),
                                 'lines': len(body.get('lines') or body.get('allocations')), 'zero_skipped': zero_skipped}
        except ValueError as e:
            tx['errors'].append(f'Baris {tx["rows"][0]}: {e}')
        txs.append(tx)
    return txs


@finance_payment_import_bp.route('', methods=['POST'])
@jwt_required()
@require_permission('finance.create')
def run_import():
    """multipart: file, type, source, commit ('1' = simpan, selain itu pratinjau)."""
    from routes.finance import create_payment_core
    try:
        ptype, source = _kind(request.form)
        f = request.files.get('file')
        if not f:
            return jsonify({'error': 'File belum dipilih'}), 400
        rows = _read_rows(f, source)
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    except Exception:
        return jsonify({'error': 'File tidak bisa dibaca. Pastikan formatnya .xlsx atau .csv sesuai template.'}), 400

    commit = request.form.get('commit') == '1'
    uid = get_jwt_identity()
    txs = build_transactions(rows, ptype, source)
    result = []
    for tx in txs:
        item = {'number': tx['number'], 'rows': tx['rows'], 'errors': list(tx['errors']), 'summary': tx.get('summary'), 'created': None}
        if tx['data'] and not item['errors']:
            try:
                resp, status = create_payment_core(tx['data'], uid, commit=commit)
            except Exception as e:                                  # one bad transaction must not stop the rest
                db.session.rollback()
                resp, status = None, 500
                item['errors'].append(str(e))
            if resp is not None:
                body = resp.get_json() or {}
                if status == 201:
                    item['created'] = body.get('payment_number')
                else:
                    db.session.rollback()
                    item['errors'].append(body.get('error') or 'Gagal menyimpan')
        result.append(item)
    if not commit:
        db.session.rollback()
    ok = [r for r in result if not r['errors']]
    return jsonify({'committed': commit, 'transactions': result,
                    'summary': {'total': len(result), 'ok': len(ok), 'error': len(result) - len(ok),
                                'amount': round(sum((r['summary'] or {}).get('total', 0) for r in ok), 2)}}), 200
