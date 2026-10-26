"""
Regression tests for the 2026-10-24 finance/accounting audit: reports and
subledgers must stay in step with the GL, and the logic gaps found while
running full flows must stay closed.
"""
from datetime import date, timedelta


ACCOUNTS = {
    '1-1010': ('Kas di Bank', 'asset', 'debit', True),
    '1-1100': ('Piutang Usaha', 'asset', 'debit', False),
    '1-2000': ('Aset Tetap', 'asset', 'debit', False),
    '1-2100': ('Akumulasi Penyusutan', 'asset', 'debit', False),  # contra, credit balance
    '2-1000': ('Hutang Usaha', 'liability', 'credit', False),
    '2-1210': ('Hutang PPN', 'liability', 'credit', False),
    '3-1000': ('Modal', 'equity', 'credit', False),
    '4-1000': ('Pendapatan Penjualan', 'revenue', 'credit', False),
    '6-6000': ('Beban Penyusutan', 'expense', 'debit', False),
    '9-1000': ('Beban Bunga', 'expense', 'debit', False),
}
_seq = {'n': 0}


def _acc(db_session, code):
    from models.finance import Account
    acc = Account.query.filter_by(account_code=code).first()
    if acc:
        return acc
    name, typ, normal, cash = ACCOUNTS[code]
    acc = Account(account_code=code, account_name=name, account_type=typ, normal_balance=normal,
                  is_active=True, is_cash_bank=cash)
    db_session.add(acc)
    db_session.commit()
    return acc


def _post(db_session, code, debit=0, credit=0, on=None, status='posted', ref_type=None, ref_id=None):
    from models.finance import AccountingEntry
    acc = _acc(db_session, code)
    _seq['n'] += 1
    db_session.add(AccountingEntry(
        entry_number=f'JE-REG-{_seq["n"]:05d}-01', entry_date=on or date.today(), entry_type='general',
        reference_type=ref_type, reference_id=ref_id,
        account_id=acc.id, account_code=code, account_name=acc.account_name,
        debit_amount=debit, credit_amount=credit, status=status,
    ))
    db_session.commit()
    return acc


# ---------------------------------------------------------------- balance sheet

def test_balance_sheet_balances_with_accumulated_depreciation(client, auth_headers, db_session):
    """Regression: accumulated depreciation (credit balance) was ADDED to fixed
    assets, so the sheet was off by exactly 2x the accumulated depreciation."""
    _post(db_session, '1-1010', debit=1_000_000)
    _post(db_session, '3-1000', credit=1_000_000)
    _post(db_session, '1-2000', debit=270_000)
    _post(db_session, '1-1010', credit=270_000)
    _post(db_session, '6-6000', debit=3_000)
    _post(db_session, '1-2100', credit=3_000)

    bs = client.get('/api/finance/reports/balance-sheet', headers=auth_headers).get_json()['balance_sheet']
    assert bs['is_balanced'], bs['imbalance']
    assert bs['assets']['fixed_assets']['net_fixed_assets'] == 267_000
    assert bs['assets']['total_assets'] == 997_000


def test_balance_sheet_cuts_off_at_report_date(client, auth_headers, db_session):
    """Regression: balances were all-time while earnings were date-filtered, so a
    sheet for a past date was out of balance whenever later transactions existed."""
    long_ago = date.today() - timedelta(days=400)
    _post(db_session, '1-1010', debit=500_000, on=long_ago)
    _post(db_session, '3-1000', credit=500_000, on=long_ago)
    _post(db_session, '1-1010', debit=200_000)              # later sale, today
    _post(db_session, '4-1000', credit=200_000)

    past = client.get(f'/api/finance/reports/balance-sheet?date={long_ago.isoformat()}', headers=auth_headers).get_json()['balance_sheet']
    assert past['is_balanced']
    assert past['assets']['total_assets'] == 500_000          # later sale not in the past sheet
    now = client.get('/api/finance/reports/balance-sheet', headers=auth_headers).get_json()['balance_sheet']
    assert now['is_balanced']
    assert now['assets']['total_assets'] == 700_000


# ------------------------------------------------------------------- cash flow

def test_cash_flow_posted_only_and_sections_add_up(client, auth_headers, db_session):
    _post(db_session, '1-1010', debit=1_000_000, ref_type='payment')
    _post(db_session, '1-1010', credit=100_000, ref_type='manual_journal')
    _post(db_session, '1-1010', credit=270_000, ref_type='asset')
    _post(db_session, '1-1010', debit=999_999, status='draft')       # must be ignored
    _post(db_session, '1-1010', debit=888_888, status='reversed')    # must be ignored

    cf = client.get(f'/api/finance/reports/cash-flow?year={date.today().year}', headers=auth_headers).get_json()['cash_flow']
    ops = cf['operating_activities']['net_cash_from_operations']
    inv = cf['investing_activities']['net_cash_from_investing']
    fin = cf['financing_activities']['net_cash_from_financing']
    assert cf['net_change_in_cash'] == 630_000
    assert ops + inv + fin == cf['net_change_in_cash']
    assert inv == -270_000
    assert cf['ending_cash'] == 630_000


# -------------------------------------------------------------- standalone cash

def _bank(db_session):
    return _acc(db_session, '1-1010')


def test_standalone_cash_cannot_post_to_receivable_or_payable_control_accounts(client, auth_headers, db_session):
    bank = _bank(db_session)
    for code in ('1-1100', '2-1000'):
        acc = _acc(db_session, code)
        resp = client.post('/api/finance/payments', headers=auth_headers, json={
            'payment_date': date.today().isoformat(), 'payment_type': 'payment', 'payment_method': 'cash',
            'bank_account_id': bank.id, 'lines': [{'account_id': acc.id, 'amount': 10}],
        })
        assert resp.status_code == 400, code
        assert 'kontrol' in resp.get_json()['error'].lower()


def test_payment_list_works_for_supplier_payments(client, auth_headers, db_session):
    """Regression: the list crashed with 500 ('Supplier' has no attribute 'name')."""
    from models.finance import Payment
    from models.purchasing import Supplier
    bank = _bank(db_session)
    sup = Supplier(code='SUP-REG', company_name='PT Pemasok Uji', is_active=True)
    db_session.add(sup)
    db_session.commit()
    db_session.add(Payment(payment_number='PAY-REG-1', payment_date=date.today(), payment_type='payment',
                           supplier_id=sup.id, payment_method='cash', amount=1000, bank_account_id=bank.id, status='pending'))
    db_session.commit()

    resp = client.get('/api/finance/payments?type=payment', headers=auth_headers)
    assert resp.status_code == 200
    assert resp.get_json()['payments'][0]['party_name'] == 'PT Pemasok Uji'


# -------------------------------------------------------------- invoice credit

def _sales_invoice(db_session, total=1_000_000, paid=0):
    from models.finance import Invoice
    inv = Invoice(invoice_number=f'INV-REG-{_seq["n"]}', invoice_type='sales', invoice_date=date.today(),
                  total_amount=total, paid_amount=paid, balance_due=total - paid, status='sent')
    _seq['n'] += 1
    db_session.add(inv)
    db_session.commit()
    return inv


def test_apply_credit_to_invoice_reduces_balance_and_survives_payment_void(db_session):
    from utils.finance_helpers import apply_credit_to_invoice
    inv = _sales_invoice(db_session, total=1_000_000)
    apply_credit_to_invoice(inv, 160_000)
    assert float(inv.balance_due) == 840_000
    assert float(inv.credit_applied) == 160_000
    assert inv.status == 'partial'

    apply_credit_to_invoice(inv, 5_000_000)           # more than is still owed
    assert float(inv.balance_due) == 0                # never negative
    assert inv.status == 'paid'


def test_void_payment_keeps_credit_applied_in_balance(client, auth_headers, db_session):
    """Voiding a payment recomputed balance as total - paid, which silently
    undid any sales-return credit already applied to the invoice."""
    from utils.finance_helpers import apply_credit_to_invoice
    from models.finance import GlobalAccountDefault
    from models.sales import Customer
    bank = _bank(db_session)
    ar = _acc(db_session, '1-1100')
    db_session.add(GlobalAccountDefault(transaction_key='accounts_receivable', account_id=ar.id))
    cust = Customer(code='CUST-REG', company_name='Pelanggan Uji', is_active=True)
    db_session.add(cust)
    db_session.commit()
    inv = _sales_invoice(db_session, total=1_000_000)
    inv.customer_id = cust.id
    db_session.commit()
    apply_credit_to_invoice(inv, 100_000)             # balance 900,000
    resp = client.post('/api/finance/payments', headers=auth_headers, json={
        'payment_date': date.today().isoformat(), 'payment_method': 'cash', 'bank_account_id': bank.id,
        'allocations': [{'invoice_id': inv.id, 'amount': 300_000}],
    })
    assert resp.status_code == 201, resp.get_json()
    pid = resp.get_json()['payment_id']
    db_session.refresh(inv)
    assert float(inv.balance_due) == 600_000

    assert client.post(f'/api/finance/payments/{pid}/void', headers=auth_headers, json={}).status_code == 200
    db_session.refresh(inv)
    assert float(inv.balance_due) == 900_000          # 1,000,000 - 100,000 credit, NOT 1,000,000


# ------------------------------------------------------------ bank reconciliation

def _statement_line(db_session, bank, amount):
    from models.finance import BankStatement, BankStatementLine
    st = BankStatement(account_id=bank.id, period_start=date.today(), period_end=date.today(),
                       source_filename='t.csv', imported_by=1)
    db_session.add(st)
    db_session.flush()
    line = BankStatementLine(statement_id=st.id, line_date=date.today(), description='x', amount=amount, is_matched=False)
    db_session.add(line)
    db_session.commit()
    return line


def test_bank_match_validates_account_amount_status_and_uniqueness(client, auth_headers, db_session):
    from models.finance import AccountingEntry
    bank = _bank(db_session)
    _post(db_session, '1-1010', debit=500_000)
    e_bank = AccountingEntry.query.filter_by(account_code='1-1010').first()
    _post(db_session, '4-1000', credit=500_000)
    e_other = AccountingEntry.query.filter_by(account_code='4-1000').first()
    _post(db_session, '1-1010', debit=500_000, status='draft')
    e_draft = AccountingEntry.query.filter_by(account_code='1-1010', status='draft').first()

    def match(line, entry):
        return client.post(f'/api/finance/bank-reconciliation/lines/{line.id}/match', headers=auth_headers,
                           json={'accounting_entry_id': entry.id})

    l1, l2, l_out = (_statement_line(db_session, bank, 500_000), _statement_line(db_session, bank, 500_000),
                     _statement_line(db_session, bank, -500_000))
    assert match(l1, e_other).status_code == 400          # wrong account
    assert match(l1, e_draft).status_code == 400          # not posted
    assert match(l_out, e_bank).status_code == 400        # wrong direction
    assert match(l1, e_bank).status_code == 200
    assert match(l2, e_bank).status_code == 400           # already matched elsewhere
    client.post(f'/api/finance/bank-reconciliation/lines/{l1.id}/unmatch', headers=auth_headers, json={})
    assert match(l2, e_bank).status_code == 200


# ------------------------------------------------------------------ depreciation

def test_monthly_depreciation_is_idempotent(db_session):
    """Regression: re-closing a reopened period crashed on the unique
    (asset, period, book) constraint - depreciation must skip assets already posted."""
    from models.asset_management import Asset, DepreciationSchedule
    from models.finance import GlobalAccountDefault
    from utils.finance_helpers import run_monthly_depreciation

    exp = _acc(db_session, '6-6000')
    acc = _acc(db_session, '1-2100')
    db_session.add_all([GlobalAccountDefault(transaction_key='depreciation_expense', account_id=exp.id),
                        GlobalAccountDefault(transaction_key='accumulated_depreciation', account_id=acc.id)])
    asset = Asset(asset_code='AST-REG', asset_name='Mesin', asset_type='machinery', status='active',
                  purchase_cost=12_000_000, salvage_value=0, useful_life_years=10,
                  depreciation_method='straight_line', accumulated_depreciation=0)
    db_session.add(asset)
    db_session.commit()

    first = run_monthly_depreciation(2026, 10, posted_by_user_id=None)
    assert first['assets_processed'] == 1
    second = run_monthly_depreciation(2026, 10, posted_by_user_id=None)
    assert second['assets_processed'] == 0 and second['total_depreciation'] == 0
    assert DepreciationSchedule.query.filter_by(asset_id=asset.id).count() == 1
    db_session.refresh(asset)
    assert float(asset.accumulated_depreciation) == 100_000


# -------------------------------------------------------------------- tax / CoA

def test_tax_page_includes_invoice_ppn_without_double_counting(client, auth_headers, db_session):
    from models.finance import Invoice, TaxTransaction
    today = date.today()
    sales = Invoice(invoice_number='INV-TAX-1', invoice_type='sales', invoice_date=today, subtotal=1_000_000,
                    tax_amount=110_000, total_amount=1_110_000, balance_due=1_110_000, status='draft')
    purchase = Invoice(invoice_number='PI-TAX-1', invoice_type='purchase', invoice_date=today, subtotal=500_000,
                       tax_amount=55_000, total_amount=555_000, balance_due=555_000, status='sent')
    cancelled = Invoice(invoice_number='INV-TAX-X', invoice_type='sales', invoice_date=today, subtotal=100,
                        tax_amount=11, total_amount=111, status='cancelled')
    db_session.add_all([sales, purchase, cancelled])
    db_session.commit()

    data = client.get('/api/finance/tax-management', headers=auth_headers).get_json()
    assert data['tax_summary']['vat_payable'] == 110_000
    assert data['tax_summary']['vat_receivable'] == 55_000
    assert {r['transaction_number'] for r in data['tax_transactions']} == {'INV-TAX-1', 'PI-TAX-1'}
    assert all(r['source'] == 'invoice' for r in data['tax_transactions'])
    # the Tax page reads type/description/tax_rate on every row (it crashed on a
    # missing 'type'), so the API must always send them
    rows = {r['transaction_number']: r for r in data['tax_transactions']}
    assert rows['INV-TAX-1']['type'] == 'VAT Output' and rows['PI-TAX-1']['type'] == 'VAT Input'
    assert all(r['description'] and 'tax_rate' in r for r in data['tax_transactions'])
    assert rows['INV-TAX-1']['tax_rate'] == 11.0

    # once the same invoice is recorded manually it must not be counted twice
    db_session.add(TaxTransaction(
        transaction_number='TX-1', transaction_date=today, transaction_type='vat_out', tax_type='PPN', tax_rate=11,
        base_amount=1_000_000, tax_amount=110_000, status='draft', reporting_period=f'{today.year}-{today.month:02d}',
        reference_type='invoice', reference_id=sales.id))
    db_session.commit()
    data = client.get('/api/finance/tax-management', headers=auth_headers).get_json()
    assert data['tax_summary']['vat_payable'] == 110_000


def test_create_account_validates_required_fields(client, auth_headers, db_session):
    assert client.post('/api/finance/chart-of-accounts', headers=auth_headers, json={'name': 'tanpa kode'}).status_code == 400
    assert client.post('/api/finance/chart-of-accounts', headers=auth_headers,
                       json={'code': '9-7777', 'name': 'X', 'type': 'bogus'}).status_code == 400
    ok = client.post('/api/finance/chart-of-accounts', headers=auth_headers, json={'code': '9-7777', 'name': 'X', 'type': 'expense'})
    assert ok.status_code == 201


# ------------------------------------------------------------------- payroll

def test_paying_payroll_posts_hutang_gaji_to_kas_and_is_idempotent(db_session):
    """Regression: marking a payroll record 'paid' posted nothing, so Hutang Gaji
    was never relieved and the bank balance ignored salaries paid out."""
    from datetime import date
    from types import SimpleNamespace
    import pytest
    from models.finance import AccountingEntry, GlobalAccountDefault
    from utils.finance_helpers import post_payroll_payment_journal, PayrollPaymentError

    ACCOUNTS['2-1110'] = ('Hutang Gaji', 'liability', 'credit', False)
    bank = _acc(db_session, '1-1010')
    hg = _acc(db_session, '2-1110')
    db_session.add_all([GlobalAccountDefault(transaction_key='cash', account_id=bank.id),
                        GlobalAccountDefault(transaction_key='payroll_net_payable', account_id=hg.id)])
    db_session.commit()

    rec = SimpleNamespace(id=777, net_salary=2_424_999)
    entries = post_payroll_payment_journal(rec, payment_date=date.today())
    assert len(entries) == 2
    debit = next(e for e in entries if float(e.debit_amount) > 0)
    credit = next(e for e in entries if float(e.credit_amount) > 0)
    assert debit.account_code == '2-1110' and float(debit.debit_amount) == 2_424_999
    assert credit.account_code == '1-1010' and float(credit.credit_amount) == 2_424_999

    assert post_payroll_payment_journal(rec, payment_date=date.today()) == []          # idempotent
    assert AccountingEntry.query.filter_by(reference_type='payroll_payment', reference_id=777).count() == 2

    with pytest.raises(PayrollPaymentError):                                           # not a Kas/Bank account
        post_payroll_payment_journal(SimpleNamespace(id=778, net_salary=1000), date.today(), bank_account_id=hg.id)


# ------------------------------------------- e-Faktur / approval / foreign currency
def test_efaktur_export_uses_dpp_after_discount_and_requires_company_npwp(client, auth_headers, db_session):
    import io
    import openpyxl
    from models.finance import Invoice
    from models.sales import Customer
    from models.settings import CompanyProfile
    cust = Customer(code='C-EF-1', company_name='PT Berpajak', tax_id='0987654321098765', billing_address='Jl. Pajak 1')
    profile = CompanyProfile.query.first() or CompanyProfile(company_name='PT Uji')
    profile.tax_id = '-'
    db_session.add_all([cust, profile])
    db_session.commit()
    db_session.add(Invoice(invoice_number='INV-EF-1', invoice_type='sales', invoice_date=date.today(), customer_id=cust.id,
                           subtotal=2_000_000, discount_amount=50_000, tax_amount=214_500, total_amount=2_164_500, status='sent'))
    db_session.commit()

    url = f'/api/finance/invoices/efaktur-export?start_date={date.today()}&end_date={date.today()}'
    assert client.get(url, headers=auth_headers).status_code == 400   # NPWP perusahaan belum valid

    profile.tax_id = '0012345678901234'
    db_session.commit()
    resp = client.get(url, headers=auth_headers)
    assert resp.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(resp.data))['DATA']
    row = next(r for r in ws.iter_rows(min_row=4, values_only=True) if r[4] == 'INV-EF-1')
    assert row[10] == 1_950_000 and row[11] == 214_500   # DPP = subtotal - diskon


def test_workflow_journal_must_balance_and_submit_works_with_string_identity(client, auth_headers, db_session):
    kas, modal = _acc(db_session, '1-1010'), _acc(db_session, '3-1000')
    line = lambda a, d, c: {'account_id': a.id, 'debit': d, 'credit': c, 'description': 'x'}
    body = lambda lines: {'transaction_type': 'stock_opname', 'transaction_id': 990001, 'journal_entry': {
        'entry_date': date.today().isoformat(), 'description': 'x', 'reference': 'x', 'lines': lines}}
    bad = client.post('/api/approval/workflows', json=body([line(kas, 5000, 0), line(modal, 0, 3000)]), headers=auth_headers)
    assert bad.status_code == 400 and 'seimbang' in bad.get_json()['error']
    ok = client.post('/api/approval/workflows', json=body([line(kas, 1000, 0), line(modal, 0, 1000)]), headers=auth_headers)
    assert ok.status_code == 201
    wf = ok.get_json()['workflow_id']
    assert client.post(f'/api/approval/workflows/{wf}/submit', json={}, headers=auth_headers).status_code == 200


def _foreign_ap(db_session, currency='USD', rate=16000, amount=1000):
    """A posted USD purchase invoice: AP carried in IDR at the invoice rate."""
    from models.finance import Invoice
    from models.purchasing import Supplier
    sup = Supplier(code=f'SUP-FX-{_seq["n"]}', company_name='PT Valas', is_active=True)
    db_session.add(sup)
    db_session.commit()
    inv = Invoice(invoice_number=f'PI-FX-{_seq["n"]}', invoice_type='purchase', invoice_date=date.today(), supplier_id=sup.id,
                  currency=currency, exchange_rate=rate, total_amount=amount, balance_due=amount, status='sent')
    _seq['n'] += 1
    db_session.add(inv)
    db_session.commit()
    ap = _post(db_session, '2-1000', credit=amount * rate)
    _post(db_session, '1-1010', debit=amount * rate)
    from models.finance import GlobalAccountDefault
    if not GlobalAccountDefault.query.filter_by(transaction_key='accounts_payable').first():
        db_session.add(GlobalAccountDefault(transaction_key='accounts_payable', account_id=ap.id))
        db_session.commit()
    return sup, inv


def _gl(db_session, code):
    from models.finance import AccountingEntry
    rows = AccountingEntry.query.filter_by(account_code=code, status='posted').all()
    return sum(float(r.credit_amount or 0) - float(r.debit_amount or 0) for r in rows)


def test_paying_foreign_currency_invoice_clears_ap_in_idr_and_books_fx_difference(client, auth_headers, db_session):
    """Regression: a USD 1,000 invoice carried at 16,000 was 'paid' by debiting AP only Rp1,000."""
    sup, inv = _foreign_ap(db_session)
    bank = _acc(db_session, '1-1010')
    ap_before = _gl(db_session, '2-1000')
    resp = client.post('/api/finance/payments', headers=auth_headers, json={
        'payment_type': 'payment', 'payment_date': date.today().isoformat(), 'supplier_id': sup.id, 'bank_account_id': bank.id,
        'payment_method': 'transfer', 'amount_idr': 16_100_000, 'allocations': [{'invoice_id': inv.id, 'amount': 1000}]})
    assert resp.status_code == 201, resp.get_json()
    assert round(_gl(db_session, '2-1000') - (ap_before - 16_000_000), 2) == 0     # AP relieved by Rp16.000.000
    assert round(_gl(db_session, '9-8000'), 2) == -100_000                          # rugi selisih kurs
    db_session.refresh(inv)
    assert inv.status == 'paid' and float(inv.balance_due) == 0

    # without amount_idr: no FX difference
    sup2, inv2 = _foreign_ap(db_session)
    resp = client.post('/api/finance/payments', headers=auth_headers, json={
        'payment_type': 'payment', 'payment_date': date.today().isoformat(), 'supplier_id': sup2.id, 'bank_account_id': bank.id,
        'payment_method': 'transfer', 'allocations': [{'invoice_id': inv2.id, 'amount': 1000}]})
    assert resp.status_code == 201
    assert round(_gl(db_session, '9-8000'), 2) == -100_000

    # IDR invoices must not accept amount_idr
    bad = client.post('/api/finance/payments', headers=auth_headers, json={
        'payment_type': 'receipt', 'payment_date': date.today().isoformat(), 'bank_account_id': bank.id,
        'payment_method': 'transfer', 'amount_idr': 5, 'lines': [{'account_id': _acc(db_session, '4-1000').id, 'amount': 5}]})
    assert bad.status_code == 400


def test_grn_requires_approved_po(client, auth_headers, db_session):
    from models.purchasing import PurchaseOrder, Supplier
    sup = Supplier(code='SUP-GRN', company_name='PT GRN', is_active=True)
    db_session.add(sup)
    db_session.commit()
    for i, status in enumerate(('draft', 'pending_approval', 'rejected', 'cancelled')):
        po = PurchaseOrder(po_number=f'PO-GRN-{i}', supplier_id=sup.id, order_date=date.today(), status=status)
        db_session.add(po)
        db_session.commit()
        resp = client.post('/api/purchasing/grn', headers=auth_headers, json={'po_id': po.id, 'supplier_id': sup.id, 'items': []})
        assert resp.status_code == 400 and 'setelah PO disetujui' in resp.get_json()['error'], status


def test_workflow_approval_refused_in_closed_period(client, auth_headers, db_session):
    from models.finance import AccountingEntry, PeriodClose
    kas, modal = _acc(db_session, '1-1010'), _acc(db_session, '3-1000')
    db_session.add(PeriodClose(period_year=2026, period_month=1))
    db_session.commit()
    line = lambda a, d, c: {'account_id': a.id, 'debit': d, 'credit': c, 'description': 'x'}
    created = client.post('/api/approval/workflows', headers=auth_headers, json={
        'transaction_type': 'stock_opname', 'transaction_id': 990002, 'journal_entry': {
            'entry_date': '2026-01-15', 'description': 'x', 'reference': 'x', 'lines': [line(kas, 900, 0), line(modal, 0, 900)]}})
    wf = created.get_json()['workflow_id']
    assert client.post(f'/api/approval/workflows/{wf}/submit', json={}, headers=auth_headers).status_code == 200
    assert client.post(f'/api/approval/workflows/{wf}/review', json={'notes': 'ok'}, headers=auth_headers).status_code == 200
    resp = client.post(f'/api/approval/workflows/{wf}/approve', json={'notes': 'ok'}, headers=auth_headers)
    assert resp.status_code == 400 and 'ditutup' in resp.get_json()['error']
    assert AccountingEntry.query.filter_by(entry_date=date(2026, 1, 15)).count() == 0


def test_nota_retur_export_lists_return_dpp_and_ppn(client, auth_headers, db_session):
    import io
    import openpyxl
    from models.finance import AccountingEntry
    from models.returns import CustomerReturn
    from models.sales import Customer
    cust = Customer(code='C-RET', company_name='PT Retur', tax_id='0987654321098765')
    db_session.add(cust)
    db_session.commit()
    ret = CustomerReturn(return_number='RET-1', customer_id=cust.id, return_date=date.today(), reason='damaged')
    db_session.add(ret)
    db_session.commit()
    acc = _acc(db_session, '4-1000')
    for n, desc, amt in ((1, 'Retur Penjualan - RET-1 (warehouse)', 1_000_000), (2, 'Koreksi PPN Keluaran - RET-1', 110_000)):
        db_session.add(AccountingEntry(entry_number=f'JE-RET-{n}-01', entry_date=date.today(), entry_type='general',
                                       reference_type='sales_return', reference_id=ret.id, account_id=acc.id, account_code='4-1000',
                                       account_name=acc.account_name, description=desc, debit_amount=amt, credit_amount=0, status='posted'))
    db_session.commit()
    resp = client.get(f'/api/finance/invoices/efaktur-retur-export?start_date={date.today()}&end_date={date.today()}', headers=auth_headers)
    assert resp.status_code == 200
    rows = list(openpyxl.load_workbook(io.BytesIO(resp.data)).active.iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == 'RET-1' and rows[0][3] == '0987654321098765' and rows[0][5] == 1_000_000 and rows[0][6] == 110_000


def test_efaktur_export_keeps_invoices_with_bad_npwp_but_warns(client, auth_headers, db_session):
    import io
    import openpyxl
    from models.finance import Invoice
    from models.sales import Customer
    from models.settings import CompanyProfile
    profile = CompanyProfile.query.first() or CompanyProfile(company_name='PT Uji')
    profile.tax_id = '0012345678901234'
    good = Customer(code='C-N1', company_name='Baik', tax_id='09.876.543.2-109.8765')
    short = Customer(code='C-N2', company_name='Pendek', tax_id='12345')
    empty = Customer(code='C-N3', company_name='Kosong', tax_id='')
    db_session.add_all([profile, good, short, empty])
    db_session.commit()
    for n, c in enumerate((good, short, empty)):
        db_session.add(Invoice(invoice_number=f'INV-NP-{n}', invoice_type='sales', invoice_date=date.today(), customer_id=c.id,
                               subtotal=1000, tax_amount=110, total_amount=1110, status='sent'))
    db_session.commit()
    resp = client.get(f'/api/finance/invoices/efaktur-export?start_date={date.today()}&end_date={date.today()}', headers=auth_headers)
    assert resp.status_code == 200
    rows = {r[4]: r for r in openpyxl.load_workbook(io.BytesIO(resp.data))['DATA'].iter_rows(min_row=4, values_only=True) if r[4]}
    assert set(rows) == {'INV-NP-0', 'INV-NP-1', 'INV-NP-2'}
    assert rows['INV-NP-0'][7] == '0987654321098765'
    assert resp.headers['X-Npwp-Warnings'] == 'INV-NP-1:format-salah,INV-NP-2:kosong'
    assert resp.headers['X-Npwp-Warning-Count'] == '2'


def test_group_chat_owner_can_update_and_delete_own_server(client, auth_headers, db_session):
    """Regression: owner checks compared int owner_id with the str JWT identity, so owners were always refused."""
    created = client.post('/api/chat/servers', json={'name': 'Tim Uji'}, headers=auth_headers)
    assert created.status_code in (200, 201), created.get_json()
    server_id = (created.get_json().get('server') or created.get_json())['id']
    assert client.put(f'/api/chat/servers/{server_id}', json={'name': 'Tim Baru'}, headers=auth_headers).status_code == 200
    assert client.delete(f'/api/chat/servers/{server_id}', headers=auth_headers).status_code == 200


def test_aging_report_buckets_and_matches_gl(client, auth_headers, db_session):
    from models.finance import Invoice
    from models.sales import Customer
    cust = Customer(code='C-AGE', company_name='PT Umur')
    db_session.add(cust)
    db_session.commit()
    today = date.today()
    for n, (days_late, amount) in enumerate(((-5, 100), (10, 200), (45, 300), (120, 400))):
        db_session.add(Invoice(invoice_number=f'INV-AGE-{n}', invoice_type='sales', invoice_date=today - timedelta(days=days_late + 30),
                               due_date=today - timedelta(days=days_late), customer_id=cust.id, total_amount=amount, balance_due=amount, status='sent'))
    db_session.add(Invoice(invoice_number='INV-AGE-PAID', invoice_type='sales', invoice_date=today, customer_id=cust.id,
                           total_amount=999, balance_due=0, status='paid'))
    db_session.commit()
    _post(db_session, '1-1100', debit=1000)
    data = client.get('/api/finance/reports/aging?type=receivable', headers=auth_headers).get_json()
    assert data['totals'] == {'current': 100, 'd1_30': 200, 'd31_60': 300, 'd61_90': 0, 'd90_plus': 400, 'total': 1000}
    assert data['gl_control_balance'] == 1000 and data['difference'] == 0
    assert data['parties'][0]['party_name'] == 'PT Umur'
    assert client.get('/api/finance/reports/aging?type=bogus', headers=auth_headers).status_code == 400


def test_bank_match_marks_payment_cleared_and_unmatch_restores(client, auth_headers, db_session):
    from models.finance import Payment, AccountingEntry
    bank = _bank(db_session)
    pay = Payment(payment_number='PAY-CLR-1', payment_date=date.today(), payment_type='receipt', payment_method='transfer',
                  amount=500_000, bank_account_id=bank.id, status='pending')
    db_session.add(pay)
    db_session.commit()
    _post(db_session, '1-1010', debit=500_000, ref_type='payment', ref_id=pay.id)
    entry = AccountingEntry.query.filter_by(reference_type='payment', reference_id=pay.id).first()
    line = _statement_line(db_session, bank, 500_000)
    base = '/api/finance/bank-reconciliation/lines'
    assert client.post(f'{base}/{line.id}/match', headers=auth_headers, json={'accounting_entry_id': entry.id}).status_code == 200
    db_session.refresh(pay)
    assert pay.status == 'cleared'
    assert client.post(f'{base}/{line.id}/unmatch', headers=auth_headers, json={}).status_code == 200
    db_session.refresh(pay)
    assert pay.status == 'pending'


def test_wip_gl_balance_by_work_order_and_receipt_credit_split(client, auth_headers, db_session, test_product, test_user):
    from models.finance import Account, GlobalAccountDefault
    from models.production import WorkOrder
    from models.material_issue import MaterialIssue
    from utils.finance_helpers import split_wip_vs_raw_material_credit
    wip = Account(account_code='1-1300', account_name='Persediaan WIP', account_type='asset', normal_balance='debit', is_active=True)
    db_session.add(wip)
    db_session.commit()
    db_session.add(GlobalAccountDefault(transaction_key='wip_persediaan', account_id=wip.id))
    wo = WorkOrder(wo_number='WO-WIP-1', product_id=test_product.id, quantity=10, uom='PCS', status='in_progress')
    db_session.add(wo)
    db_session.commit()
    mi = MaterialIssue(issue_number='MI-WIP-1', work_order_id=wo.id, status='issued', requested_by=test_user.id)
    db_session.add(mi)
    db_session.commit()
    _post(db_session, '1-1300', debit=1000, ref_type='material_issue', ref_id=mi.id)

    rows = client.get('/api/wip-accounting/gl-balance', headers=auth_headers).get_json()
    assert rows['gl_balance'] == 1000 and rows['rows'][0]['wo_number'] == 'WO-WIP-1' and rows['rows'][0]['balance'] == 1000

    assert split_wip_vs_raw_material_credit(wo.id, 1000) == (1000.0, 0.0)       # all via WIP
    assert split_wip_vs_raw_material_credit(wo.id, 1500) == (1000.0, 500.0)     # rest from raw material
    _post(db_session, '1-1300', credit=400, ref_type='work_order', ref_id=wo.id)  # earlier receipt took 400
    assert split_wip_vs_raw_material_credit(wo.id, 1000) == (600.0, 400.0)


def _customer_and_ar(db_session, name='PT Kredit'):
    from models.finance import GlobalAccountDefault
    from models.sales import Customer
    cust = Customer(code=f'C-CM-{_seq["n"]}', company_name=name)
    _seq['n'] += 1
    db_session.add(cust)
    ar = _acc(db_session, '1-1100')
    if not GlobalAccountDefault.query.filter_by(transaction_key='accounts_receivable').first():
        db_session.add(GlobalAccountDefault(transaction_key='accounts_receivable', account_id=ar.id))
    db_session.commit()
    return cust


def test_return_credit_beyond_invoice_becomes_credit_memo_then_apply_and_refund(client, auth_headers, db_session):
    from models.finance import CreditMemo, Invoice
    from utils.finance_helpers import credit_customer_for_return
    cust = _customer_and_ar(db_session)
    bank = _acc(db_session, '1-1010')
    inv = Invoice(invoice_number='INV-CM-1', invoice_type='sales', invoice_date=date.today(), customer_id=cust.id,
                  total_amount=1000, balance_due=300, paid_amount=700, status='partial')
    other = Invoice(invoice_number='INV-CM-2', invoice_type='sales', invoice_date=date.today(), customer_id=cust.id,
                    total_amount=500, balance_due=500, status='sent')
    db_session.add_all([inv, other])
    db_session.commit()
    _post(db_session, '1-1100', credit=800)          # the return's AR credit (already posted)

    applied, memo = credit_customer_for_return(inv, cust.id, 800, 99, date.today(), None)
    db_session.commit()
    assert applied == 300 and memo is not None and float(memo.amount) == 500        # invoice absorbs only what it owes
    db_session.refresh(inv)
    assert float(inv.balance_due) == 0 and inv.status == 'paid'

    # no linked invoice -> whole amount becomes a memo
    _, memo2 = credit_customer_for_return(None, cust.id, 100, 98, date.today(), None)
    db_session.commit()
    assert float(memo2.amount) == 100

    # apply 200 of memo 1 to the other invoice; over-application and wrong customer are refused
    url = f'/api/finance/credit-memos/{memo.id}'
    assert client.post(f'{url}/apply', headers=auth_headers, json={'invoice_id': other.id, 'amount': 600}).status_code == 400
    assert client.post(f'{url}/apply', headers=auth_headers, json={'invoice_id': other.id, 'amount': 200}).status_code == 200
    db_session.refresh(other)
    assert float(other.balance_due) == 300 and float(other.credit_applied) == 200

    # refund the remaining 300: Dr AR / Cr Kas posted, memo closed
    before = _gl(db_session, '1-1010')
    resp = client.post(f'{url}/refund', headers=auth_headers, json={'bank_account_id': bank.id, 'amount': 300})
    assert resp.status_code == 200, resp.get_json()
    assert round(_gl(db_session, '1-1010') - before, 2) == 300   # _gl = credit - debit: Kas credited 300
    db_session.refresh(memo)
    assert memo.status == 'closed' and memo.balance == 0
    assert client.post(f'{url}/refund', headers=auth_headers, json={'bank_account_id': bank.id, 'amount': 1}).status_code == 400

    listing = client.get('/api/finance/credit-memos?status=open', headers=auth_headers).get_json()
    assert [m['memo_number'] for m in listing['credit_memos']] == [memo2.memo_number] and listing['open_balance'] == 100
    aging = client.get('/api/finance/reports/aging?type=receivable', headers=auth_headers).get_json()
    assert aging['credit_memo_balance'] == 100


def test_credit_memo_applications_can_be_undone(client, auth_headers, db_session):
    from models.finance import CreditMemo, Invoice
    from utils.finance_helpers import credit_customer_for_return
    cust = _customer_and_ar(db_session, 'PT Batal')
    bank = _acc(db_session, '1-1010')
    inv = Invoice(invoice_number='INV-UN-1', invoice_type='sales', invoice_date=date.today(), customer_id=cust.id,
                  total_amount=500, balance_due=500, status='sent')
    db_session.add(inv)
    db_session.commit()
    _post(db_session, '1-1100', credit=400)
    _, memo = credit_customer_for_return(None, cust.id, 400, 77, date.today(), None)
    db_session.commit()
    url = f'/api/finance/credit-memos/{memo.id}'

    assert client.post(f'{url}/apply', headers=auth_headers, json={'invoice_id': inv.id, 'amount': 150}).status_code == 200
    assert client.post(f'{url}/refund', headers=auth_headers, json={'bank_account_id': bank.id, 'amount': 250}).status_code == 200
    detail = client.get(url, headers=auth_headers).get_json()
    apply_id, refund_id = [a['id'] for a in detail['applications']]
    assert detail['balance'] == 0
    cash_before = _gl(db_session, '1-1010')

    # undo refund: cash comes back, memo balance restored, second undo refused
    assert client.post(f'/api/finance/credit-memos/applications/{refund_id}/undo', headers=auth_headers).status_code == 200
    assert round(_gl(db_session, '1-1010') - cash_before, 2) == -250       # reversing journal debits Kas back
    assert client.post(f'/api/finance/credit-memos/applications/{refund_id}/undo', headers=auth_headers).status_code == 400
    db_session.refresh(memo)
    assert memo.balance == 250

    # undo invoice use: invoice owes it again
    assert client.post(f'/api/finance/credit-memos/applications/{apply_id}/undo', headers=auth_headers).status_code == 200
    db_session.refresh(inv)
    db_session.refresh(memo)
    assert float(inv.balance_due) == 500 and float(inv.credit_applied) == 0 and inv.status == 'sent'
    assert memo.balance == 400 and memo.status == 'open'
    assert [a['reversed'] for a in client.get(url, headers=auth_headers).get_json()['applications']] == [True, True]


def test_manual_stock_in_posts_inventory_and_opening_balance_journal(client, auth_headers, db_session, test_material, test_product):
    """Regression: stock entered by hand never touched the GL, so shipping it later credited inventory for goods
    that were never debited (inventory went negative)."""
    from models.finance import AccountingEntry, GlobalAccountDefault
    from models.warehouse import WarehouseZone, WarehouseLocation
    from models.finance import Account
    inv_acc = Account(account_code='1-1200', account_name='Persediaan Bahan Baku', account_type='asset', normal_balance='debit', is_active=True)
    db_session.add(inv_acc)
    db_session.commit()
    db_session.add(GlobalAccountDefault(transaction_key='akun_persediaan_id', account_id=inv_acc.id))
    test_material.cost_per_unit = 2000
    zone = WarehouseZone(code='ZONE-MS', name='Zone MS', material_type='raw_materials', zone_type='storage')
    db_session.add(zone)
    db_session.commit()
    loc = WarehouseLocation(zone_id=zone.id, location_code='ZONE-MS-R1-L1-P1', rack='R1', level='L1', position='P1', capacity_uom='PCS')
    db_session.add(loc)
    db_session.commit()

    def balance(code):
        rows = AccountingEntry.query.filter_by(account_code=code, status='posted').all()
        return sum(float(r.debit_amount or 0) - float(r.credit_amount or 0) for r in rows)

    r = client.post('/api/warehouse/inventory/add', headers=auth_headers,
                    json={'material_id': test_material.id, 'location_id': loc.id, 'quantity': 10, 'uom': 'PCS'})
    assert r.status_code in (200, 201), r.get_json()
    assert balance('1-1200') == 20_000 and balance('3-4000') == -20_000

    r = client.post('/api/warehouse/stock-input/quick-add', headers=auth_headers,
                    json={'material_id': test_material.id, 'location_id': loc.id, 'quantity': 5})
    assert r.status_code in (200, 201), r.get_json()
    assert balance('1-1200') == 30_000

    # purchase-linked stock-in (GRN already journals it) must NOT be journaled again
    client.post('/api/warehouse/inventory/add', headers=auth_headers,
                json={'material_id': test_material.id, 'location_id': loc.id, 'quantity': 1, 'uom': 'PCS', 'reference_type': 'grn'})
    assert balance('1-1200') == 30_000

    # item without master cost -> no journal, no failure
    test_material.cost_per_unit = 0
    db_session.commit()
    r = client.post('/api/warehouse/stock-input/quick-add', headers=auth_headers,
                    json={'material_id': test_material.id, 'location_id': loc.id, 'quantity': 3})
    assert r.status_code in (200, 201) and balance('1-1200') == 30_000


def test_balance_sheet_counts_every_equity_account(client, auth_headers, db_session):
    """Regression: only 3-1000/3-2000/3-3000 were summed, so Saldo Awal Persediaan (3-4000) unbalanced the sheet."""
    from models.finance import Account
    db_session.add(Account(account_code='3-4000', account_name='Saldo Awal Persediaan', account_type='equity',
                           normal_balance='credit', is_active=True))
    db_session.commit()
    _post(db_session, '1-1010', debit=1_000)
    _post(db_session, '3-4000', credit=1_000)
    data = client.get('/api/finance/reports/balance-sheet', headers=auth_headers).get_json()['balance_sheet']
    assert data['is_balanced'] is True and data['equity']['total_equity'] == 1_000
