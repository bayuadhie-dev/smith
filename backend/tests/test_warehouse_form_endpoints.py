import io
import openpyxl
from datetime import datetime, timedelta


def _loc(db_session, code='LOC-T1'):
    from models import WarehouseLocation, WarehouseZone
    z = WarehouseZone.query.first()
    if not z:
        z = WarehouseZone(code='ZT', name='Zona Tes', material_type='finished_goods', zone_type='storage', is_active=True)
        db_session.add(z); db_session.commit()
    loc = WarehouseLocation(location_code=code, zone_id=z.id, rack='R1', level='1', position='1', capacity=1000, capacity_uom='PCS', is_active=True)
    db_session.add(loc); db_session.commit()
    return loc


def _inv(db_session, product, loc, qty=100, reserved=10):
    from models import Inventory
    inv = Inventory(product_id=product.id, location_id=loc.id, quantity_on_hand=qty, quantity_reserved=reserved,
                    quantity_available=qty - reserved, batch_number='B1', is_active=True)
    db_session.add(inv); db_session.commit()
    return inv


def test_inventory_get_put_available_and_ledger(client, auth_headers, db_session, test_product):
    from models.wms_advanced import InventoryTransaction
    loc = _loc(db_session)
    inv = _inv(db_session, test_product, loc)
    d = client.get(f'/api/warehouse/inventory/{inv.id}', headers=auth_headers).get_json()
    assert d['quantity'] == 100 and d['reserved_quantity'] == 10 and d['batch_number'] == 'B1'
    av = client.get(f'/api/warehouse/inventory/available?product_id={test_product.id}&location_id={loc.id}', headers=auth_headers).get_json()
    assert av['available_quantity'] == 90
    before = InventoryTransaction.query.count()
    r = client.put(f'/api/warehouse/inventory/{inv.id}', headers=auth_headers, json={'quantity': 120, 'reserved_quantity': 20, 'lot_number': 'L9'})
    assert r.status_code == 200, r.get_json()
    d = client.get(f'/api/warehouse/inventory/{inv.id}', headers=auth_headers).get_json()
    assert d['quantity'] == 120 and d['available_quantity'] == 100 and d['lot_number'] == 'L9'
    t = InventoryTransaction.query.order_by(InventoryTransaction.id.desc()).first()
    assert InventoryTransaction.query.count() == before + 1 and t.direction == 'in' and float(t.quantity) == 20
    assert client.put(f'/api/warehouse/inventory/{inv.id}', headers=auth_headers, json={'quantity': 5, 'reserved_quantity': 9}).status_code == 400
    assert client.get('/api/warehouse/inventory/999999', headers=auth_headers).status_code == 404


def test_transfer_request_crud_and_validation(client, auth_headers, db_session, test_product):
    a, b = _loc(db_session, 'LOC-A'), _loc(db_session, 'LOC-B')
    _inv(db_session, test_product, a, qty=50, reserved=0)
    body = {'product_id': test_product.id, 'from_location_id': a.id, 'to_location_id': b.id, 'quantity': 20, 'uom': 'PCS', 'priority': 'high'}
    r = client.post('/api/warehouse/transfers', headers=auth_headers, json=body)
    assert r.status_code == 201, r.get_json()
    tid = r.get_json()['id']
    assert client.post('/api/warehouse/transfers', headers=auth_headers, json={**body, 'quantity': 500}).status_code == 400
    assert client.post('/api/warehouse/transfers', headers=auth_headers, json={**body, 'to_location_id': a.id}).status_code == 400
    assert client.put(f'/api/warehouse/transfers/{tid}', headers=auth_headers, json={**body, 'quantity': 30}).status_code == 200
    g = client.get(f'/api/warehouse/transfers/{tid}', headers=auth_headers).get_json()
    assert g['quantity'] == 30 and g['priority'] == 'high' and g['status'] == 'pending'
    assert client.get('/api/warehouse/transfers', headers=auth_headers).get_json()['total'] == 1


def test_movements_export_is_xlsx(client, auth_headers, db_session):
    from utils.inventory_helpers import record_inventory_transaction
    record_inventory_transaction(transaction_type='stock_in', direction='in', quantity=5, notes='ekspor tes')
    db_session.commit()
    r = client.get('/api/warehouse/movements/export?search=ekspor', headers=auth_headers)
    assert r.status_code == 200 and r.data[:2] == b'PK'
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    assert ws['A1'].value == 'No. Mutasi' and ws.max_row == 2


def test_schedule_get_put(client, auth_headers, db_session):
    from models.production import ProductionSchedule, WorkOrder
    from models import Machine
    from models import Product
    p = Product(code='PSCH', name='P Sched', primary_uom='PCS', price=0, cost=0, material_type='finished_goods')
    db_session.add(p); db_session.commit()
    wo = WorkOrder(wo_number='WO-SCH-T', product_id=p.id, quantity=10, uom='PCS', status='planned')
    m = Machine(code='MC-SCH', name='Mesin Sched', machine_type='nonwoven_machine')
    db_session.add_all([wo, m]); db_session.commit()
    s = ProductionSchedule(schedule_number='SCH-T1', work_order_id=wo.id, machine_id=m.id, scheduled_start=datetime(2026, 11, 1, 8), scheduled_end=datetime(2026, 11, 1, 16), status='scheduled')
    db_session.add(s); db_session.commit()
    assert client.get(f'/api/production/schedules/{s.id}', headers=auth_headers).get_json()['schedule_number'] == 'SCH-T1'
    r = client.put(f'/api/production/schedules/{s.id}', headers=auth_headers, json={'scheduled_start': '2026-11-02T08:00', 'scheduled_end': '2026-11-02T18:00', 'shift': 'shift_2'})
    assert r.status_code == 200 and r.get_json()['shift'] == 'shift_2'
    assert client.put(f'/api/production/schedules/{s.id}', headers=auth_headers, json={'scheduled_start': '2026-11-02T08:00', 'scheduled_end': '2026-11-02T07:00'}).status_code == 400
