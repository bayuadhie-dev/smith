"""Finance report center (routes/finance_reports.py)."""
import io
from datetime import date, timedelta

import openpyxl

from models import db
from models.finance import Account, AccountingEntry, Invoice, Payment, PaymentAllocation

_n = {'i': 0}


def _acc(code, name, typ, normal, header=False, cash=False, parent=None):
    a = Account.query.filter_by(account_code=code).first()
    if a:
        return a
    a = Account(account_code=code, account_name=name, account_type=typ, normal_balance=normal, is_header=header,
                is_cash_bank=cash, is_active=True, parent_id=parent.id if parent else None)
    db.session.add(a)
    db.session.commit()
    return a


def _je(acc, debit=0, credit=0, on=None, ref_type=None, ref_id=None, status='posted'):
    _n['i'] += 1
    db.session.add(AccountingEntry(entry_number=f'JE-RPT-{_n["i"]:05d}-01', entry_date=on or date.today(), entry_type='general',
                                   account_id=acc.id, account_code=acc.account_code, account_name=acc.account_name,
                                   debit_amount=debit, credit_amount=credit, reference_type=ref_type, reference_id=ref_id, status=status))
    db.session.commit()


def _world():
    bank_h = _acc('1001', 'Bank', 'asset', 'debit', header=True)
    bca = _acc('100101', 'BCA', 'asset', 'debit', cash=True, parent=bank_h)
    ina = _acc('100106', 'INA', 'asset', 'debit', cash=True, parent=bank_h)
    ar = _acc('100201', 'Piutang', 'asset', 'debit')
    modal = _acc('300003', 'Modal', 'equity', 'credit')
    sales = _acc('400002', 'Penjualan', 'revenue', 'credit')
    cogs = _acc('500001', 'HPP', 'expense', 'debit')
    d1, d2, d3 = date.today() - timedelta(days=40), date.today() - timedelta(days=5), date.today()
    _je(bca, debit=1000, on=d1); _je(modal, credit=1000, on=d1)               # capital before the period
    _je(bca, debit=300, on=d2); _je(sales, credit=300, on=d2)                  # receipt in the period
    _je(bca, credit=100, on=d3); _je(cogs, debit=100, on=d3)                   # payment today
    _je(ina, debit=50, on=d3); _je(sales, credit=50, on=d3)
    _je(bca, debit=999, on=d3, status='draft')                                  # draft: must not count
    return dict(bank_h=bank_h, bca=bca, ina=ina, ar=ar, modal=modal, sales=sales, cogs=cogs, d1=d1, d2=d2, d3=d3)


def _q(client, auth_headers, url):
    return client.get(url, headers=auth_headers)


def test_trial_balance_balances_and_uses_opening_and_period(client, auth_headers, db_session):
    w = _world()
    start = (w['d3'] - timedelta(days=10)).isoformat()
    data = _q(client, auth_headers, f'/api/finance/reports/trial-balance?date_from={start}&date_to={w["d3"]}').get_json()
    assert data['is_balanced'] is True
    rows = {r['account_code']: r for r in data['rows']}
    assert rows['100101']['opening_debit'] == 1000 and rows['100101']['period_debit'] == 300 and rows['100101']['period_credit'] == 100
    assert rows['100101']['closing_debit'] == 1200            # draft 999 ignored
    assert rows['300003']['closing_credit'] == 1000 and '1001' not in rows   # header accounts are not listed
    assert data['totals']['closing_debit'] == data['totals']['closing_credit']


def test_account_ledger_running_balance_and_header_aggregation(client, auth_headers, db_session):
    w = _world()
    url = f'/api/finance/reports/account-ledger?account_id={w["bca"].id}&date_from={(w["d3"] - timedelta(days=10)).isoformat()}&date_to={w["d3"]}'
    data = _q(client, auth_headers, url).get_json()
    assert data['opening_balance'] == 1000
    assert [r['balance'] for r in data['rows']] == [1300, 1200]
    assert data['total_debit'] == 300 and data['total_credit'] == 100 and data['closing_balance'] == 1200
    hdr = _q(client, auth_headers, f'/api/finance/reports/account-ledger?account_id={w["bank_h"].id}&date_from=2000-01-01&date_to={w["d3"]}').get_json()
    assert hdr['closing_balance'] == 1250 and {r['account_code'] for r in hdr['rows']} == {'100101', '100106'}   # BCA 1200 + INA 50
    assert _q(client, auth_headers, '/api/finance/reports/account-ledger?account_id=999999').status_code == 404


def test_cash_daily_and_position(client, auth_headers, db_session):
    w = _world()
    start = (w['d3'] - timedelta(days=10)).isoformat()
    data = _q(client, auth_headers, f'/api/finance/reports/cash-daily?date_from={start}&date_to={w["d3"]}').get_json()
    assert data['totals'] == {'opening': 1000.0, 'receipts': 350.0, 'disbursements': 100.0, 'closing': 1250.0}
    assert [(r['date'], r['receipts'], r['disbursements'], r['closing']) for r in data['rows']] == [
        (w['d2'].isoformat(), 300.0, 0.0, 1300.0), (w['d3'].isoformat(), 50.0, 100.0, 1250.0)]
    only = _q(client, auth_headers, f'/api/finance/reports/cash-daily?date_from={start}&date_to={w["d3"]}&account_id={w["ina"].id}').get_json()
    assert only['totals']['closing'] == 50.0
    pos = _q(client, auth_headers, f'/api/finance/reports/cash-position?as_of={w["d3"]}').get_json()
    assert pos['total'] == 1250.0 and {r['code']: r['balance'] for r in pos['rows']} == {'100101': 1200.0, '100106': 50.0}


def test_reports_export_to_excel(client, auth_headers, db_session):
    w = _world()
    resp = _q(client, auth_headers, f'/api/finance/reports/trial-balance?date_from=2000-01-01&date_to={w["d3"]}&format=xlsx')
    assert resp.status_code == 200 and 'spreadsheetml' in resp.mimetype
    ws = openpyxl.load_workbook(io.BytesIO(resp.data)).active
    assert ws['A2'].value == 'Neraca Saldo' and ws['A6'].value == 'Kode' and ws.freeze_panes == 'A7'
    last = [c.value for c in ws[ws.max_row]]
    assert last[1] == 'TOTAL' and ws.cell(row=ws.max_row, column=4).number_format.startswith('#,##0')
    assert _q(client, auth_headers, f'/api/finance/reports/account-ledger?account_id={w["bca"].id}&format=xlsx').status_code == 200
    assert _q(client, auth_headers, '/api/finance/reports/trial-balance?date_from=2026-02-01&date_to=2026-01-01').status_code == 400


def _customer_world():
    from models.sales import Customer
    cust = Customer(code='C-RPT', company_name='PT Pelanggan', tax_id='0987654321098765')
    db.session.add(cust)
    db.session.commit()
    d = date.today()
    inv1 = Invoice(invoice_number='INV-R1', invoice_type='sales', invoice_date=d - timedelta(days=40), customer_id=cust.id, subtotal=1000, tax_amount=110,
                   total_amount=1110, balance_due=1110, status='sent')
    inv2 = Invoice(invoice_number='INV-R2', invoice_type='sales', invoice_date=d - timedelta(days=2), customer_id=cust.id, subtotal=500, tax_amount=55,
                   total_amount=555, balance_due=555, status='sent')
    voided = Invoice(invoice_number='INV-RX', invoice_type='sales', invoice_date=d, customer_id=cust.id, total_amount=9999, status='cancelled')
    db.session.add_all([inv1, inv2, voided])
    db.session.commit()
    bank = _acc('100101', 'BCA', 'asset', 'debit', cash=True)
    pay = Payment(payment_number='RCP-R1', payment_date=d - timedelta(days=1), payment_type='receipt', customer_id=cust.id, payment_method='transfer',
                  amount=600, bank_account_id=bank.id, status='pending')
    gone = Payment(payment_number='RCP-RX', payment_date=d, payment_type='receipt', customer_id=cust.id, payment_method='transfer', amount=777,
                   bank_account_id=bank.id, status='cancelled')
    db.session.add_all([pay, gone])
    db.session.commit()
    return cust


def test_customer_statement_and_summary_agree(client, auth_headers, db_session):
    cust = _customer_world()
    start = (date.today() - timedelta(days=10)).isoformat()
    s = _q(client, auth_headers, f'/api/finance/reports/party-statement/customer/{cust.id}?date_from={start}&date_to={date.today()}').get_json()
    assert s['opening_balance'] == 1110                      # INV-R1 before the period
    assert [r['number'] for r in s['rows']] == ['INV-R2', 'RCP-R1']   # cancelled invoice / voided payment excluded
    assert [r['balance'] for r in s['rows']] == [1665, 1065] and s['closing_balance'] == 1065
    summary = _q(client, auth_headers, f'/api/finance/reports/party-statements?type=customer&date_from={start}&date_to={date.today()}').get_json()
    row = summary['rows'][0]
    assert (row['opening'], row['invoiced'], row['payments'], row['closing']) == (1110, 555, 600, 1065)
    assert summary['totals']['closing'] == 1065
    assert _q(client, auth_headers, f'/api/finance/reports/party-statement/customer/{cust.id}?format=xlsx').status_code == 200
    assert _q(client, auth_headers, '/api/finance/reports/party-statement/bogus/1').status_code == 400


def test_supplier_statement_converts_foreign_invoices_at_invoice_rate(client, auth_headers, db_session):
    from models.purchasing import Supplier
    sup = Supplier(code='S-RPT', company_name='PT Pemasok', is_active=True)
    db.session.add(sup)
    db.session.commit()
    d = date.today()
    inv = Invoice(invoice_number='PI-R1', invoice_type='purchase', invoice_date=d, supplier_id=sup.id, currency='USD', exchange_rate=16000, total_amount=100,
                  subtotal=100, balance_due=100, status='sent')
    db.session.add(inv)
    db.session.commit()
    bank = _acc('100101', 'BCA', 'asset', 'debit', cash=True)
    pay = Payment(payment_number='PAY-R1', payment_date=d, payment_type='payment', supplier_id=sup.id, payment_method='transfer', amount=40, currency='USD', bank_account_id=bank.id, status='pending')
    db.session.add(pay)
    db.session.commit()
    db.session.add(PaymentAllocation(payment_id=pay.id, invoice_id=inv.id, amount=40))
    db.session.commit()
    s = _q(client, auth_headers, f'/api/finance/reports/party-statement/supplier/{sup.id}?date_from={d}&date_to={d}').get_json()
    assert [(r['number'], r['amount']) for r in s['rows']] == [('PI-R1', 1_600_000.0), ('PAY-R1', -640_000.0)]
    assert s['closing_balance'] == 960_000.0


def test_purchase_report_vat_and_equity(client, auth_headers, db_session):
    from models.purchasing import Supplier
    sup = Supplier(code='S-RPT2', company_name='PT Pemasok 2', tax_id='01.234', is_active=True)
    db.session.add(sup)
    db.session.commit()
    d = date.today()
    db.session.add_all([
        Invoice(invoice_number='PI-V1', invoice_type='purchase', invoice_date=d, supplier_id=sup.id, subtotal=1000, tax_amount=110, total_amount=1110, balance_due=1110, status='sent'),
        Invoice(invoice_number='INV-V1', invoice_type='sales', invoice_date=d, subtotal=2000, tax_amount=220, total_amount=2220, balance_due=2220, status='sent'),
    ])
    db.session.commit()
    pr = _q(client, auth_headers, f'/api/finance/reports/purchase-report?date_from={d}&date_to={d}').get_json()
    assert pr['totals']['total'] == 1110 and pr['by_supplier'][0]['supplier_name'] == 'PT Pemasok 2' and pr['rows'][0]['po_number'] is None
    vat = _q(client, auth_headers, f'/api/finance/reports/vat-summary?year={d.year}').get_json()
    this = next(r for r in vat['rows'] if r['month'] == f'{d.year}-{d.month:02d}')
    assert (this['output_vat'], this['input_vat'], this['net_payable']) == (220, 110, 110)
    det = _q(client, auth_headers, f'/api/finance/reports/vat-detail?month={d.year}-{d.month:02d}&kind=input').get_json()
    assert det['rows'][0]['tax_id'] == '01.234' and det['totals']['vat'] == 110
    assert _q(client, auth_headers, '/api/finance/reports/vat-detail?month=bad').status_code == 400

    w = _world()
    eq = _q(client, auth_headers, f'/api/finance/reports/equity-changes?date_from={(w["d3"] - timedelta(days=10)).isoformat()}&date_to={w["d3"]}').get_json()
    assert eq['rows'][0]['code'] == '300003' and eq['rows'][0]['opening'] == 1000 and eq['period_profit'] == 250 and eq['total_closing'] == 1250


def test_cash_daily_detail_lists_lines_with_counter_accounts(client, auth_headers, db_session):
    w = _world()
    data = _q(client, auth_headers, f'/api/finance/reports/cash-daily-detail?date={w["d3"]}').get_json()
    assert data['totals'] == {'receipts': 50.0, 'disbursements': 100.0}
    assert {r['account_code'] for r in data['rows']} == {'100101', '100106'}
    assert _q(client, auth_headers, '/api/finance/reports/cash-daily-detail').status_code == 400


def test_every_report_exports_to_excel_and_real_pdf(client, auth_headers, db_session):
    w = _world()
    cust = _customer_world()
    rng = f'date_from=2000-01-01&date_to={w["d3"]}'
    month = w['d3'].strftime('%Y-%m')
    urls = [f'trial-balance?{rng}', f'account-ledger?account_id={w["bca"].id}&{rng}', f'cash-daily?{rng}', f'cash-daily-detail?date={w["d3"]}',
            f'cash-position?date_to={w["d3"]}', f'party-statements?type=customer&{rng}', f'party-statement/customer/{cust.id}?{rng}',
            f'purchase-report?{rng}', f'purchase-report?{rng}&group=supplier', f'vat-summary?year={w["d3"].year}',
            f'vat-detail?month={month}&kind=output', f'equity-changes?{rng}']
    for u in urls:
        for fmt, marker in (('xlsx', b'PK'), ('pdf', b'%PDF')):
            resp = _q(client, auth_headers, f'/api/finance/reports/{u}&format={fmt}')
            assert resp.status_code == 200, (u, fmt, resp.data[:200])
            assert resp.data[:4].startswith(marker), (u, fmt)
