"""Impor Penerimaan/Pembayaran dari Excel/CSV (routes/finance_payment_import.py)."""
import io
from datetime import date

import openpyxl

from tests.test_finance_sync_regressions import _acc, _bank, _sales_invoice


def _xlsx(rows, headers):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Transaksi'
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


KAS = ['Tanggal', 'No. Bukti', 'Kas/Bank', 'Akun Lawan', 'Nilai', 'Pihak', 'Keterangan', 'No. Referensi']


def _post(client, auth_headers, buf, ptype='payment', source='kas', commit=False, name='x.xlsx'):
    return client.post('/api/finance/payment-import', headers=auth_headers, content_type='multipart/form-data',
                       data={'file': (buf, name), 'type': ptype, 'source': source, 'commit': '1' if commit else '0'})


def test_template_downloads_for_every_kind(client, auth_headers, db_session):
    for ptype in ('receipt', 'payment'):
        for source in ('kas', 'invoice'):
            r = client.get(f'/api/finance/payment-import/template?type={ptype}&source={source}', headers=auth_headers)
            assert r.status_code == 200 and r.data[:2] == b'PK'
            ws = openpyxl.load_workbook(io.BytesIO(r.data))['Transaksi']
            assert ws['A1'].value == 'Tanggal' and ws.max_row >= 2
    assert client.get('/api/finance/payment-import/template?type=x&source=kas', headers=auth_headers).status_code == 400


def test_preview_validates_without_saving_and_commit_saves(client, auth_headers, db_session):
    from models.finance import AccountingEntry, Payment
    bank = _bank(db_session)
    beban, potongan = _acc(db_session, '6-6000'), _acc(db_session, '9-1000')
    today = date.today().isoformat()
    rows = [
        [today, '', bank.account_code, beban.account_code, 100000, 'PT A', 'Biaya', 'REF1'],
        [today, 'IMP-2', bank.account_code, beban.account_code, 100, '', 'Baris 1 bukti dua', ''],
        [today, 'IMP-2', bank.account_code, potongan.account_code, -30, '', 'Pengurang', ''],
        [today, '', bank.account_code, 'TIDAK-ADA', 5000, '', '', ''],
        [today, '', bank.account_code, beban.account_code, 0, '', '', ''],
        ['bukan-tanggal', '', bank.account_code, beban.account_code, 5, '', '', ''],
    ]
    before = Payment.query.count()
    prev = _post(client, auth_headers, _xlsx(rows, KAS)).get_json()
    assert prev['committed'] is False and prev['summary'] == {'total': 5, 'ok': 2, 'error': 3, 'amount': 100070.0}
    assert Payment.query.count() == before                         # pratinjau tidak menyimpan
    errors = [e for t in prev['transactions'] for e in t['errors']]
    assert any('TIDAK-ADA' in e for e in errors) and any('0' in e for e in errors) and any('bukan-tanggal' in e for e in errors)

    done = _post(client, auth_headers, _xlsx(rows, KAS), commit=True).get_json()
    assert done['committed'] is True and done['summary']['ok'] == 2 and done['summary']['error'] == 3
    created = {t['number'] or t['created']: t['created'] for t in done['transactions'] if t['created']}
    assert 'IMP-2' in created and Payment.query.count() == before + 2
    pay = Payment.query.filter_by(payment_number='IMP-2').first()
    assert float(pay.amount) == 70 and len(pay.lines) == 2
    es = AccountingEntry.query.filter_by(reference_type='payment', reference_id=pay.id, status='posted').all()
    assert round(sum(float(e.debit_amount) for e in es) - sum(float(e.credit_amount) for e in es), 2) == 0
    dup = _post(client, auth_headers, _xlsx(rows[1:3], KAS), commit=True).get_json()
    assert dup['summary']['ok'] == 0 and 'sudah dipakai' in dup['transactions'][0]['errors'][0]


def test_invoice_import_reports_wrong_invoice_and_posts_valid_one(client, auth_headers, db_session):
    from models.finance import GlobalAccountDefault
    from models.sales import Customer
    bank = _bank(db_session)
    db_session.add(GlobalAccountDefault(transaction_key='accounts_receivable', account_id=_acc(db_session, '1-1100').id))
    cust = Customer(code='C-IMP', company_name='PT Impor', is_active=True)
    db_session.add(cust)
    db_session.commit()
    inv = _sales_invoice(db_session, total=1_000_000)
    inv.customer_id = cust.id
    db_session.commit()
    today = date.today().isoformat()
    head = ['Tanggal', 'No. Bukti', 'Kas/Bank', 'No. Faktur', 'Nilai', 'Keterangan', 'No. Referensi']
    rows = [[today, 'IMP-INV', bank.account_code, inv.invoice_number, 400000, 'Cicilan', 'TRF'],
            [today, '', bank.account_code, 'INV-TIDAK-ADA', 1000, '', ''],
            [today, '', bank.account_code, inv.invoice_number, 9_999_999, '', '']]
    r = _post(client, auth_headers, _xlsx(rows, head), ptype='receipt', source='invoice', commit=True).get_json()
    assert r['summary']['ok'] == 1 and r['summary']['error'] == 2
    db_session.refresh(inv)
    assert float(inv.paid_amount) == 400_000 and inv.status == 'partial'
    wrong_kind = _post(client, auth_headers, _xlsx(rows[:1], head), ptype='payment', source='invoice').get_json()
    assert 'bukan faktur pembelian' in wrong_kind['transactions'][0]['errors'][0]


def test_bad_files_are_rejected_clearly(client, auth_headers, db_session):
    assert _post(client, auth_headers, io.BytesIO(b'x'), name='x.txt').status_code == 400
    assert _post(client, auth_headers, _xlsx([], KAS)).status_code == 400                      # no data rows
    bad_head = _post(client, auth_headers, _xlsx([[1, 2]], ['A', 'B']))
    assert bad_head.status_code == 400 and 'Kolom wajib' in bad_head.get_json()['error']
    csv_buf = io.BytesIO('Tanggal;Kas/Bank;Akun Lawan;Nilai\n2026-01-01;1-1010;9-1000;1.500,50\n'.encode())
    _bank(db_session); _acc(db_session, '9-1000')
    r = _post(client, auth_headers, csv_buf, name='x.csv').get_json()
    assert r['summary']['ok'] == 1 and r['transactions'][0]['summary']['total'] == 1500.5


def test_zero_value_lines_are_skipped_not_rejected(client, auth_headers, db_session):
    from models.finance import Payment
    bank = _bank(db_session)
    beban, lain = _acc(db_session, '6-6000'), _acc(db_session, '9-1000')
    today = date.today().isoformat()
    rows = [
        [today, 'IMP-Z', bank.account_code, beban.account_code, 50000, 'PT Z', 'Baris isi', ''],
        [today, 'IMP-Z', bank.account_code, lain.account_code, 0, '', 'Baris nol', ''],
        [today, 'IMP-ALLZERO', bank.account_code, beban.account_code, 0, '', '', ''],
    ]
    res = _post(client, auth_headers, _xlsx(rows, KAS), commit=True).get_json()
    assert res['summary']['ok'] == 1 and res['summary']['error'] == 1
    ok = [t for t in res['transactions'] if not t['errors']][0]
    assert ok['summary']['lines'] == 1 and ok['summary']['zero_skipped'] == 1
    assert Payment.query.filter_by(payment_number='IMP-Z').count() == 1
    assert any('bernilai 0' in e for t in res['transactions'] for e in t['errors'])
